# -*- coding: utf-8 -*-
"""按 JSON 规格填充买来的 PPT 模板。

用法：
    tpl-fill template.pptx out.pptx spec.json
    # 未安装时：cd tools/pptx-fill && python -m tplkit.fill_deck template.pptx out.pptx spec.json

规格（spec.json）结构 —— 键是【模板里的原始页号，1 起】：

{
  "final_order": ["1","3","9","8","12","16","10","13","15","17"],
  "delete_slides": [2,4,5,6,7,11,14,18],
  "slides": {
    "1": {
      "text":   {"T0": "新文字", "T2": "多行\\n第二行"},
      "widenT": {"T0": 6.3},                       // 改宽，保持水平居中
      "boxT":   {"T0": {"left":1.0,"top":2.0,"width":6.3,"height":0.6}},  // 显式设框（可只写几个键），优先级高于 widenT
      "hideP":  [0,1],                             // 按图片序号删除
      "hideT":  [3],                               // 按文本框序号删除
      "hideName": ["组合 9"],                       // 按形状名删除
      "pic":    {"P1": "path/to/img.png"},          // 替换图片/图片填充
      "addpic": [{"path":"...","left":1.0,"top":2.0,"width":6.0}],
      "addtext":[{"text":"...","left":1.0,"top":2.0,"width":6.0,"height":0.6,"size":18,"bold":true,"color":"FFFFFF","align":"center"}]
    }
  }
}

序号规则与 dump-spec.py 完全一致（文档顺序，图片含图片填充的矩形），见 tplkit/shapes.py。
"""
import json
import sys
from copy import deepcopy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Emu, Inches, Pt

from tplkit.shapes import A_NS, R_NS, index_map, qn


def set_text(shape, text):
    """保留首个 run 的格式，替换文字。支持 \\n 分段。"""
    tf = shape.text_frame
    body = tf._txBody
    ps = body.findall(qn("a:p"))
    if not ps:
        return
    p0 = ps[0]
    runs = p0.findall(qn("a:r"))
    if not runs:
        return
    r0 = runs[0]
    for r in runs[1:]:
        p0.remove(r)
    for p in ps[1:]:
        body.remove(p)
    t = r0.find(qn("a:t"))
    if t is None:
        t = r0.makeelement(qn("a:t"), {})
        r0.append(t)
    lines = str(text).split("\n")
    t.text = lines[0]
    prev = p0
    for line in lines[1:]:
        np = deepcopy(p0)
        nr = np.findall(qn("a:r"))[0]
        nt = nr.find(qn("a:t"))
        if nt is None:
            nt = nr.makeelement(qn("a:t"), {})
            nr.append(nt)
        nt.text = line
        prev.addnext(np)
        prev = np


def drop_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def replace_pic(slide, shape, path):
    blips = shape._element.findall(f".//{{{A_NS}}}blip")
    if not blips:
        return False
    blip = blips[0]
    old_rid = blip.get(f"{{{R_NS}}}embed")
    _, new_rid = slide.part.get_or_add_image_part(path)
    blip.set(f"{{{R_NS}}}embed", new_rid)
    # 旧图片若已无人引用，断开关系，避免包体膨胀
    if old_rid and old_rid != new_rid:
        xml = slide._element.xml
        if f'r:embed="{old_rid}"' not in xml and f"r:link=\"{old_rid}\"" not in xml:
            try:
                slide.part.drop_rel(old_rid)
            except Exception:
                pass
    return True


def main():
    tpl, out, spec_path = sys.argv[1], sys.argv[2], sys.argv[3]
    spec = json.loads(open(spec_path, encoding="utf-8").read())
    prs = Presentation(tpl)
    slides = list(prs.slides)

    for key, page in spec.get("slides", {}).items():
        idx = int(key) - 1
        if idx < 0 or idx >= len(slides):
            print(f"  !! 页号 {key} 不存在，跳过")
            continue
        slide = slides[idx]
        tmap, pmap, acc = index_map(slide)
        n_t, n_p = len(tmap), len(pmap)

        for t, txt in page.get("text", {}).items():
            i = int(t[1:])
            if i not in tmap:
                print(f"  !! slide {key} 没有 {t}（该页共 {n_t} 个文本框）")
                continue
            set_text(tmap[i], txt)

        for t, w in page.get("widenT", {}).items():
            i = int(t[1:])
            if i in tmap:
                sh = tmap[i]
                cx = sh.left + sh.width / 2
                sh.width = Inches(w)
                sh.left = int(cx - sh.width / 2)

        for t, box in page.get("boxT", {}).items():
            i = int(t[1:])
            if i not in tmap:
                print(f"  !! slide {key} 没有 {t}（该页共 {n_t} 个文本框）")
                continue
            sh = tmap[i]
            if "left" in box:
                sh.left = Inches(box["left"])
            if "top" in box:
                sh.top = Inches(box["top"])
            if "width" in box:
                sh.width = Inches(box["width"])
            if "height" in box:
                sh.height = Inches(box["height"])

        for t, sz in page.get("sizeT", {}).items():
            i = int(t[1:])
            if i in tmap:
                for para in tmap[i].text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(sz)

        for i in page.get("hideT", []):
            if i in tmap:
                drop_shape(tmap[i])
        for i in page.get("hideP", []):
            if i in pmap:
                drop_shape(pmap[i])
        for nm in page.get("hideName", []):
            hit = [s for k, s in acc if s.name == nm]
            for s in hit:
                drop_shape(s)
            if not hit:
                print(f"  !! slide {key} 找不到形状名 {nm}")

        for p, path in page.get("pic", {}).items():
            i = int(p[1:])
            if i not in pmap:
                print(f"  !! slide {key} 没有 {p}（该页共 {n_p} 张图）")
                continue
            replace_pic(slide, pmap[i], path)

        for a in page.get("addpic", []):
            pic = slide.shapes.add_picture(
                a["path"], Inches(a["left"]), Inches(a["top"]),
                width=Inches(a["width"]) if "width" in a else None,
                height=Inches(a["height"]) if "height" in a else None,
            )
            if a.get("shadow"):
                pic._element.spPr.append(parse_xml(
                    f"<a:effectLst {nsdecls('a')}><a:outerShdw blurRad='203200'"
                    " dist='38100' dir='5400000' rotWithShape='0'>"
                    "<a:srgbClr val='000000'><a:alpha val='20000'/></a:srgbClr>"
                    "</a:outerShdw></a:effectLst>"))

        for r in page.get("addrect", []):
            shp = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(r["left"]), Inches(r["top"]),
                Inches(r["width"]), Inches(r["height"]))
            shp.line.fill.background()
            shp.shadow.inherit = False
            fill = shp.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string(r.get("color", "000000"))
            solid = fill._xPr.find(qn("a:solidFill"))
            clr = solid.find(qn("a:srgbClr"))
            if clr is not None:
                clr.append(parse_xml(f"<a:alpha {nsdecls('a')} val='{int(r.get('alpha', 0.5) * 100000)}'/>"))

        for a in page.get("addtext", []):
            box = slide.shapes.add_textbox(Inches(a["left"]), Inches(a["top"]),
                                           Inches(a["width"]), Inches(a.get("height", 0.6)))
            tf = box.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            run = para.add_run()
            run.text = a["text"]
            if a.get("align") in ("left", "center", "right"):
                para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                  "right": PP_ALIGN.RIGHT}[a["align"]]
            if "size" in a:
                run.font.size = Pt(a["size"])
            if a.get("bold"):
                run.font.bold = True
            if a.get("color"):
                run.font.color.rgb = RGBColor.from_string(a["color"])
            if a.get("glow"):
                run.font.color.rgb = RGBColor.from_string(a.get("color", "FFFFFF"))
                run._r.get_or_add_rPr().append(parse_xml(
                    f"<a:effectLst {nsdecls('a')}><a:outerShdw blurRad='63500'"
                    " dist='0' dir='0' algn='ctr' rotWithShape='0'>"
                    "<a:srgbClr val='000000'><a:alpha val='65000'/></a:srgbClr>"
                    "</a:outerShdw></a:effectLst>"))

    # 删页 + 重排
    sldIdLst = prs.slides._sldIdLst
    elems = list(sldIdLst)
    keep = []
    for i, el in enumerate(elems, 1):
        if i in spec.get("delete_slides", []):
            rid = el.get(qn("r:id"))
            prs.part.drop_rel(rid)
            sldIdLst.remove(el)
        else:
            keep.append((i, el))

    order = spec.get("final_order")
    if order:
        by_num = {n: el for n, el in keep}
        seq = [by_num[int(n)] for n in order if int(n) in by_num]
        for el in elems:
            if el.getparent() is not None:
                sldIdLst.remove(el)
        for el in seq:
            sldIdLst.append(el)

    prs.save(out)
    print(f"已写出 {out}（{len(prs.slides.__iter__.__self__._sldIdLst)} 页）")


if __name__ == "__main__":
    main()