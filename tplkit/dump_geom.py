# -*- coding: utf-8 -*-
"""dump 成品 pptx：每页每个文本框的编号(T#)/位置/尺寸/字号/文字，图片的位置尺寸。

用法：
    tpl-geom tmp/outputs/deck-filled-v6.pptx
编号规则与 tplkit/shapes.py 一致：文档顺序，只数「文本非空的形状」和图片。
"""
import sys

from pptx import Presentation

from tplkit.shapes import EMU_IN, collect


def sizes(sh):
    out = []
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                out.append(r.font.size.pt)
    return sorted(set(out))


def main():
    prs = Presentation(sys.argv[1])
    print(f"{sys.argv[1]}: {prs.slide_width/EMU_IN:.2f}x{prs.slide_height/EMU_IN:.2f}in")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n===== final p{i} =====")
        ti = pi = 0
        for kind, sh in collect(slide.shapes):
            try:
                pos = f"({sh.left/EMU_IN:5.2f},{sh.top/EMU_IN:5.2f}) {sh.width/EMU_IN:5.2f}x{sh.height/EMU_IN:5.2f}"
            except TypeError:
                pos = "?"
            if kind == "TXT":
                txt = "\n".join(p.text for p in sh.text_frame.paragraphs).replace("\n", " ⏎ ")
                print(f" T{ti} [{sh.name}] {pos}  pt={sizes(sh)}")
                print(f"      {txt}")
                ti += 1
            elif kind == "PIC":
                print(f" P{pi} [PIC {sh.name}] {pos}")
                pi += 1


if __name__ == "__main__":
    main()