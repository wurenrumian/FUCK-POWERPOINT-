# -*- coding: utf-8 -*-
"""把模板每一页的形状结构 dump 出来，用于决定「哪一页填什么」。

用法：
    tpl-dump-slides template.pptx [--pages 1,3,9] [--full]

输出：每个形状的名字 / 类型 / 位置尺寸（英寸）/ 文字（截断）。
--full 不截断文字。
"""
import sys

from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0


def walk(shapes, depth=0, out=None):
    """递归展开组合。"""
    pad = "  " * depth
    for sh in shapes:
        st = str(sh.shape_type).replace("AUTO_SHAPE_TYPE.", "").replace("TEXT_BOX", "TEXTBOX")
        try:
            pos = f"({sh.left/EMU_IN:5.2f},{sh.top/EMU_IN:5.2f}) {sh.width/EMU_IN:5.2f}x{sh.height/EMU_IN:5.2f}in"
        except TypeError:
            pos = "(no pos)"

        if sh.shape_type is not None and "GROUP" in str(sh.shape_type):
            out.append(f"{pad}[G] {sh.name} {pos}")
            walk(sh.shapes, depth + 1, out)
            continue

        txt = ""
        has_tf = getattr(sh, "has_text_frame", False)
        if has_tf and sh.has_text_frame:
            paras = [p.text.strip() for p in sh.text_frame.paragraphs]
            paras = [p for p in paras if p]
            txt = " | ".join(paras)
        elif sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append(" / ".join(c.text.strip() for c in r.cells))
            txt = " <TABLE> ".join(rows)
        elif sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
            try:
                img = sh.image
                txt = f"<PIC {img.filename or img.content_type} {len(img.blob)//1024}KB>"
            except Exception:
                txt = "<PIC>"
        out.append(f"{pad}[{st}] {sh.name} {pos}")
        if txt:
            out.append(f"{pad}     “{txt}”")


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return
    path = args[0]
    pages = None
    full = "--full" in args
    if "--pages" in args:
        pages = {int(x) for x in args[args.index("--pages") + 1].split(",")}

    prs = Presentation(path)
    print(f"{path}: {len(prs.slides)} slides, {prs.slide_width/EMU_IN:.2f} x {prs.slide_height/EMU_IN:.2f} in")

    for i, slide in enumerate(prs.slides, 1):
        if pages and i not in pages:
            continue
        layout = slide.slide_layout.name if slide.slide_layout is not None else "?"
        print(f"\n{'='*78}\n幻灯片 {i}   (layout: {layout})\n{'='*78}")
        out = []
        walk(slide.shapes, 0, out)
        for line in out:
            if not full and "“" in line and len(line) > 220:
                line = line[:220] + "…”"
            print(line)


if __name__ == "__main__":
    main()