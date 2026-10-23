# -*- coding: utf-8 -*-
"""详细 dump：形状名 / 类型 / 位置尺寸 / 段落对齐 / 垂直锚点 / 自动缩放 / 每段字号与字体。

用法：
    tpl-dump-detail template.pptx --pages 1,8,10,16,17
编号沿用 tplkit/shapes.py 的规则（文档顺序，只数文本非空形状 + 图片），另外把 OTHER 也列出来。
"""
import sys

from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

from tplkit.shapes import EMU_IN, collect


def body_pr(sh):
    try:
        return sh.text_frame._txBody.find(qn("a:bodyPr"))
    except Exception:
        return None


def describe_tf(sh):
    tf = sh.text_frame
    bp = body_pr(sh)
    info = []
    if bp is not None:
        for tag, key in (("a:normAutofit", "normAutofit"), ("a:spAutoFit", "spAutoFit"),
                         ("a:noAutofit", "noAutofit")):
            if bp.find(qn(tag)) is not None:
                info.append(key)
        info.append("anchor=" + (bp.get("anchor") or "-"))
        info.append("wrap=" + (bp.get("wrap") or "1"))
        l, r, t, b = bp.get("lIns"), bp.get("rIns"), bp.get("tIns"), bp.get("bIns")
        info.append(f"ins=({l},{t},{r},{b})")
    lines = []
    for p in tf.paragraphs:
        if not p.text.strip():
            continue
        al = p.alignment
        szs = []
        for r in p.runs:
            f = r.font
            szs.append(f"{f.size.pt if f.size else '-'}/{f.bold and 'B' or '-'}/{f.name or r.font._rPr.get('latin', {}).get('typeface') if False else (f.name or '-')}")
        lines.append(f"      [{al}] " + " ".join(szs) + f"  «{p.text}»")
    return info, lines


def main():
    args = sys.argv[1:]
    path = args[0]
    pages = None
    if "--pages" in args:
        pages = {int(x) for x in args[args.index("--pages") + 1].split(",")}
    prs = Presentation(path)
    print(f"{path}: {prs.slide_width/EMU_IN:.2f}x{prs.slide_height/EMU_IN:.2f}in, {len(prs.slides)} slides")
    for i, slide in enumerate(prs.slides, 1):
        if pages and i not in pages:
            continue
        print(f"\n===== template slide {i} =====")
        ti = pi = oi = 0
        for kind, sh in collect(slide.shapes):
            try:
                pos = f"({sh.left/EMU_IN:5.2f},{sh.top/EMU_IN:5.2f}) {sh.width/EMU_IN:5.2f}x{sh.height/EMU_IN:5.2f}"
            except TypeError:
                pos = "?"
            if kind == "TXT":
                info, lines = describe_tf(sh)
                print(f" T{ti} [{sh.name}] {pos}  {' '.join(info)}")
                for l in lines:
                    print(l)
                ti += 1
            elif kind == "PIC":
                print(f" P{pi} [PIC {sh.name}] {pos}")
                pi += 1
            else:
                print(f" O{oi} [OTHER {sh.name}] {pos}  {sh.shape_type}")
                oi += 1


if __name__ == "__main__":
    main()