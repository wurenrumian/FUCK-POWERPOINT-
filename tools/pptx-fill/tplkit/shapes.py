# -*- coding: utf-8 -*-
"""形状索引的共用实现。

这里是 `T#/P#` 编号规则的**唯一定义**：`fill_deck`（定位并修改）、`dump_geom` /
`dump_detail`（回读成品）都 import 本模块，`dump_spec`（生成骨架）沿用同一套规则。
"改了其中一边，另一边必须一起改"这条约定，从这个模块起算。

规则：按**文档顺序**遍历；递归进入组合；
只把「文字非空的形状」计为 T#、「图片（含矩形上的图片填充 blipFill）」计为 P#。
"""

from __future__ import annotations

EMU_IN = 914400.0
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def qn(tag):
    from pptx.oxml.ns import qn as _qn
    return _qn(tag)


def is_group(sh):
    return sh.shape_type is not None and "GROUP" in str(sh.shape_type)


def collect(shapes, acc=None, include_groups=False):
    """返回 [(kind, shape)]，kind in {TXT, PIC, OTHER, GRP}，文档顺序，进组合。

    include_groups=True 时把「组合」本身也记进 acc（kind="GRP"）——`fill_deck` 的
    `hideName` 需要按形状名找到组合；其它调用方不需要，行为与从前一致。
    """
    acc = [] if acc is None else acc
    for sh in shapes:
        if is_group(sh):
            if include_groups:
                acc.append(("GRP", sh))
            collect(sh.shapes, acc, include_groups)
            continue
        kind = str(sh.shape_type)
        has_blip = sh._element.find(f".//{{{A_NS}}}blip") is not None
        if "PICTURE" in kind:
            acc.append(("PIC", sh))
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and any(
            p.text.strip() for p in sh.text_frame.paragraphs
        ):
            acc.append(("TXT", sh))
        elif has_blip:
            acc.append(("PIC", sh))
        else:
            acc.append(("OTHER", sh))
    return acc


def index_map(slide):
    """返回 (tmap, pmap, acc)：T# 形状表、P# 形状表、完整索引（含 GRP，供 hideName 查找）。"""
    acc = collect(slide.shapes, include_groups=True)
    tmap, pmap = {}, {}
    ti = pi = 0
    for kind, sh in acc:
        if kind == "TXT":
            tmap[ti] = sh
            ti += 1
        elif kind == "PIC":
            pmap[pi] = sh
            pi += 1
    return tmap, pmap, acc